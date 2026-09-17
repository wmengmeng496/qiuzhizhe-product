from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path('后端三个源码文件.docx')
FILES = [
    Path('/Users/wuqiong/Desktop/后端文件/docx-reader.ts'),
    Path('/Users/wuqiong/Desktop/后端文件/page.tsx'),
    Path('/Users/wuqiong/Desktop/后端文件/pdf-reader.ts'),
]

def set_font(run, name='Menlo', size=8.5, color='202124', bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold

def shade(paragraph, fill='F5F7FA'):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); ppr.append(shd)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.65); sec.right_margin = Inches(0.65)
sec.header_distance = Inches(0.3); sec.footer_distance = Inches(0.3)

normal = doc.styles['Normal']
normal.font.name = 'Calibri'; normal.font.size = Pt(10)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

code_style = doc.styles.add_style('Source Code', WD_STYLE_TYPE.PARAGRAPH)
code_style.font.name = 'Menlo'; code_style.font.size = Pt(8.5); code_style.font.color.rgb = RGBColor(32,33,36)
code_style._element.rPr.rFonts.set(qn('w:ascii'), 'Menlo'); code_style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Menlo')
code_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Menlo')
code_style.paragraph_format.space_before = Pt(0); code_style.paragraph_format.space_after = Pt(0)
code_style.paragraph_format.line_spacing = 1.0

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(6)
r = title.add_run('后端三个源码文件'); set_font(r, 'Calibri', 22, '17365D', True)
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER; sub.paragraph_format.space_after = Pt(16)
r = sub.add_run('完整代码归档 · 保留原始缩进、空行与文件顺序'); set_font(r, 'Calibri', 10, '5F6368')

for idx, path in enumerate(FILES):
    if idx > 0: doc.add_page_break()
    h = doc.add_paragraph(); h.paragraph_format.space_before = Pt(4); h.paragraph_format.space_after = Pt(8)
    r = h.add_run(path.name); set_font(r, 'Calibri', 16, '2E74B5', True)
    meta = doc.add_paragraph(); meta.paragraph_format.space_after = Pt(6)
    r = meta.add_run(f'路径：{path}\n行数：{len(path.read_text(encoding="utf-8").splitlines())}'); set_font(r, 'Calibri', 9, '5F6368')
    for n, line in enumerate(path.read_text(encoding='utf-8').splitlines(), 1):
        p = doc.add_paragraph(style='Source Code')
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.keep_together = True
        shade(p, 'F8F9FA' if n % 2 == 0 else 'FFFFFF')
        rr = p.add_run(f'{n:04d}  {line}')
        set_font(rr)

header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rh = header.add_run('后端源码归档'); set_font(rh, 'Calibri', 8, '7A7F87')
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
rf = footer.add_run('代码内容按源文件原样收录'); set_font(rf, 'Calibri', 8, '7A7F87')

doc.save(OUT)
print(OUT.resolve())
