from pathlib import Path
import re
from docx import Document

SOURCE_DOC = Path('/Users/wuqiong/Desktop/后端代码源码文件.docx')
OUTPUT_DOC = Path('后端代码源码文件_去掉序号.docx')
SOURCE_FILES = {
    'docx-reader.ts': Path('src/lib/docx-reader.ts'),
    'page.tsx': Path('src/pages/home/page.tsx'),
    'pdf-reader.ts': Path('src/lib/pdf-reader.ts'),
}

doc = Document(SOURCE_DOC)
paragraphs = doc.paragraphs

for filename, source_path in SOURCE_FILES.items():
    header_index = next(
        i for i, paragraph in enumerate(paragraphs)
        if (paragraph.text.splitlines() and paragraph.text.splitlines()[0] == filename)
    )
    lines = source_path.read_text(encoding='utf-8').splitlines()
    start = header_index + 1
    if start < len(paragraphs) and paragraphs[start].text.startswith('行数：'):
        paragraphs[start].runs[0].text = f'行数：{len(lines)}'
        start += 1
    elif start < len(paragraphs) and '\n行数：' in paragraphs[start].text:
        start += 1

    code_paragraphs = []
    pos = start
    while pos < len(paragraphs):
        text = paragraphs[pos].text
        first_line = text.splitlines()[0] if text.splitlines() else ''
        if pos > start and first_line in SOURCE_FILES:
            break
        if re.match(r'^\d{4}  ', text) or paragraphs[pos].style.name == 'Source Code':
            code_paragraphs.append(paragraphs[pos])
        pos += 1

    if len(code_paragraphs) != len(lines):
        raise RuntimeError(f'{filename}: document has {len(code_paragraphs)} code lines, source has {len(lines)}')

    for paragraph, line in zip(code_paragraphs, lines):
        if paragraph.runs:
            paragraph.runs[0].text = line
            for run in paragraph.runs[1:]:
                run.text = ''
        else:
            paragraph.add_run(line)

doc.save(OUTPUT_DOC)
print(OUTPUT_DOC.resolve())
