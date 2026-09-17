from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('/Users/wuqiong/Desktop/文件夹们/我自己的/求职者产品/职引简历AI求职工作台系统_操作说明书_V1.1_人话版.docx')
MEDIA = Path('/private/tmp/zhiyin_doc_media')

BLUE = '2457A6'
INK = '182230'
MUTED = '5B6B7F'
LIGHT = 'E8EEF5'
CALLOUT = 'F4F7FB'

def set_font(run, size=11, color=INK, bold=False, italic=False):
    run.font.name = 'Calibri'
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Calibri')
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def cell_margin(cell, top=90, start=130, bottom=90, end=130):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_table_widths(table, widths):
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    total = sum(widths)
    tblW.set(qn('w:w'), str(total))
    tblW.set(qn('w:type'), 'dxa')
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(width))
            tcW.set(qn('w:type'), 'dxa')
            cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_para(doc, text='', size=11, color=INK, bold=False, italic=False, align=None, before=0, after=6, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    if text:
        set_font(p.add_run(text), size, color, bold, italic)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text), 11)
    return p

def add_step(doc, n, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(f'{n}. '), 11, BLUE, True)
    set_font(p.add_run(title + '：'), 11, INK, True)
    set_font(p.add_run(body), 11)
    return p

def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, CALLOUT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(label + '  '), 10, BLUE, True)
    set_font(p.add_run(text), 10, MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, widths)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, LIGHT)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(h), 10, INK, True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(value), 10, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), {1:16, 2:13, 3:12}[level], BLUE if level < 3 else '1F4D78', True)
    return p

def add_screenshot(doc, name, caption, note):
    path = MEDIA / name
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run()
        run.add_picture(str(path), width=Inches(6.45))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(3)
        set_font(cap.add_run(caption), 9, MUTED, True)
        add_para(doc, note, 9.5, MUTED, italic=True, after=8)

def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)

def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(0.78); section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.92); section.right_margin = Inches(0.92)
    section.header_distance = Inches(0.35); section.footer_distance = Inches(0.35)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    for level, size, color, before, after in [(1,16,BLUE,18,10),(2,13,BLUE,14,7),(3,12,'1F4D78',10,5)]:
        st = doc.styles[f'Heading {level}']
        st.font.name = 'Calibri'; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
        st._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
        st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after); st.paragraph_format.keep_with_next = True
    for style_name in ['List Bullet', 'List Number']:
        st = doc.styles[style_name]; st.font.name = 'Calibri'; st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run('职引简历｜AI求职工作台操作说明书'), 8.5, MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run('V1.1  ·  '), 8.5, MUTED)
    add_page_number(footer)

    # Cover / masthead
    add_para(doc, '产品操作说明书', 11, BLUE, True, align=WD_ALIGN_PARAGRAPH.CENTER, before=36, after=14)
    add_para(doc, '职引简历', 29, INK, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, 'AI求职工作台系统', 19, BLUE, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    add_para(doc, '把岗位要求、简历内容和面试准备放在一个地方完成。', 12, MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=34)
    add_table(doc, ['项目', '内容'], [
        ('软件名称', '职引简历AI求职工作台系统'),
        ('版本', 'V1.1（功能与说明书同步版）'),
        ('适用人群', '需要改简历、准备面试、整理求职资料的求职者'),
        ('使用方式', '打开浏览器访问系统，不需要安装客户端'),
        ('编制日期', '2026年8月17日'),
    ], [2200, 7160])
    add_callout(doc, '先看这里', '这份说明书按真实使用顺序来写：先登录，再建目标岗位，接着整理简历、生成优化结果，最后管理历史版本、准备面试和完善职业形象。')
    add_para(doc, '说明：原说明书中的截图保留作为界面示例；由于产品持续迭代，个别按钮名称、数据和颜色可能随线上版本更新。', 9.5, MUTED, italic=True, after=12)
    doc.add_page_break()

    add_heading(doc, '一、先用一句话了解产品', 1)
    add_para(doc, '职引简历不是“输入一句话就结束”的聊天工具，而是一张围绕求职过程搭起来的工作台。你可以先告诉它想投什么岗位，再放入自己的简历，系统会帮你找出岗位关键词、重写重点经历，并把不同岗位的版本保存下来。投递之后，还能继续用面试助手练题、做模拟面试和记录复盘。')
    add_table(doc, ['模块', '它解决什么问题', '你会在哪里使用'], [
        ('求职工作台', '把目标 JD、简历和优化指令放在同一个任务里处理。', '左侧“求职工作台”'),
        ('面试助手', '练习高频题、模拟回答、查看技巧、记录面试复盘。', '左侧“面试助手”'),
        ('优化记录', '找回不同公司、不同岗位的历史简历版本。', '左侧“优化记录”'),
        ('职业形象', '生成证件照风格图片、整理名片和对外展示资料。', '左侧“职业形象”'),
        ('积分账单', '查看积分余额、获得来源和扣除明细。', '左侧“积分账单”'),
        ('会员额度', '查看套餐、权益和开通申请状态。', '左侧“会员额度”'),
        ('个人中心', '管理账号资料、登录方式和基础设置。', '头像或“个人中心”'),
    ], [1700, 4860, 2800])
    add_callout(doc, '推荐顺序', '第一次使用时，按“求职工作台 → 优化记录 → 面试助手 → 职业形象 → 积分账单/会员额度”的顺序走，最容易形成完整闭环。')

    add_heading(doc, '二、开始使用：打开、登录和注册', 1)
    add_heading(doc, '2.1 打开系统', 2)
    add_para(doc, '在 Chrome、Edge、Safari 等现代浏览器中打开系统网址即可。产品是 Web 应用，不需要下载客户端，也不需要自己安装 Node.js、数据库或其他开发环境。电脑、平板和手机浏览器都可以访问；这份说明书以桌面浏览器为例。')
    add_heading(doc, '2.2 登录', 2)
    add_step(doc, 1, '进入登录页', '在首页点击“登录”，或在需要保存记录时按页面提示进入登录。')
    add_step(doc, 2, '选择登录方式', '当前版本支持手机号或邮箱登录；根据页面显示填写账号和密码。')
    add_step(doc, 3, '提交登录', '点击“登录”。如果账号或密码不对，页面会告诉你重新检查，不会清空已填写的其他内容。')
    add_heading(doc, '2.3 注册', 2)
    add_para(doc, '没有账号时，点击“立即注册”。手机号注册需要填写短信验证码；邮箱注册需要填写邮箱、密码和确认密码。密码至少 6 位。注册成功后会自动进入原来想去的页面，例如你从会员页发起登录，登录后会回到会员页。')
    add_screenshot(doc, 'image2.png', '图1  登录页示例（原版截图）', '截图展示的是旧版邮箱登录样式；当前版本已在同一位置增加手机号/邮箱切换和手机号验证码注册。')

    add_heading(doc, '三、求职工作台：一次优化从这里开始', 1)
    add_para(doc, '进入工作台后，页面通常分成三块：左侧是功能导航，中间是对话/任务列表，右侧是当前岗位和简历操作区。你不需要一次把所有信息填完，先输入 JD，再逐步补充简历和要求即可。')
    add_screenshot(doc, 'image1.png', '图2  求职工作台示例（原版截图）', '界面中的“AI 求职材料工作台”是旧版副标题，当前产品已统一更新为“AI求职工作台”。')
    add_heading(doc, '3.1 新建一个岗位任务', 2)
    add_step(doc, 1, '点击“新对话”', '准备投递新公司或新岗位时，新建一个对话，避免不同岗位的关键词混在一起。')
    add_step(doc, 2, '确认目标岗位', '在“目标 JD”卡片里查看岗位名称和关键词。已有岗位可以点“更新 JD”修改。')
    add_step(doc, 3, '保存岗位上下文', '确认后，后面的简历优化都会围绕这个岗位进行。换岗位时，建议重新建对话。')
    add_heading(doc, '3.2 粘贴 JD 或读取岗位链接', 2)
    add_para(doc, '最稳妥的方式是直接把招聘页面里的岗位名称、职责、任职要求、技术栈和加分项复制到“粘贴 JD”输入框。也可以选择“BOSS 链接/读取岗位链接”，按页面提示提供链接或岗位信息。系统会把内容整理成岗位名称、关键词和匹配重点，方便你先确认方向是否读对。')
    add_callout(doc, '小建议', 'JD 不要只粘贴一句岗位名称。职责和任职要求越完整，系统越容易判断哪些经历应该放前面、哪些关键词值得补进简历。')

    add_heading(doc, '四、放入简历：上传、粘贴和检查', 1)
    add_para(doc, '确认目标岗位后，再放入简历。系统支持上传 PDF、DOCX、TXT，也支持直接粘贴文字。对于扫描版 PDF，如果里面没有可复制的文字层，系统可能读不出内容；这时请改用 DOCX/TXT，或者把简历正文粘贴进输入框。')
    add_step(doc, 1, '选择文件或粘贴文字', '点击上传入口选择本地简历，或把简历正文粘贴到输入区。')
    add_step(doc, 2, '等待解析', '系统会提取姓名、简介、技能、工作经历、项目经历、教育经历和作品链接等常见模块。')
    add_step(doc, 3, '快速检查', '先看姓名、公司、时间、数字和技术栈有没有读错。AI 只能帮你改表达，不能替你确认事实。')
    add_step(doc, 4, '补充缺失内容', '如果某段经历没有被识别，直接粘贴缺失段落并说明“这是我的项目经历”即可。')
    add_callout(doc, '隐私提醒', '上传前建议先删掉身份证号、家庭住址等与求职无关的敏感信息；手机号和邮箱只保留你愿意出现在简历上的版本。')

    add_heading(doc, '五、AI 优化：让结果更像你的简历', 1)
    add_para(doc, '系统支持“整体诊断”和按模块处理。整体诊断适合第一次整理一份完整简历；模块优化适合只改某一段，不需要整份重写。当前工作台常见模块包括开场定位、技能匹配、工作经历、项目亮点、教育补强、证书资质和作品材料。')
    add_table(doc, ['你想做什么', '可以直接这样说'], [
        ('整体优化', '“基于当前 JD，把整份简历改成适合投递的版本。”'),
        ('改项目经历', '“只优化项目经历，写清楚我做了什么、怎么做、带来什么结果。”'),
        ('改自我介绍', '“把开场介绍改得更像面试时能说出口的版本，控制在 60 秒。”'),
        ('控制真实性', '“只使用我提供的事实，不要新增没有发生过的数字。”'),
        ('调整风格', '“更简洁一点，保留技术细节，但不要写成营销文案。”'),
    ], [2300, 7060])
    add_heading(doc, '5.1 看懂一次优化结果', 2)
    add_para(doc, '生成结果通常会包含优化后的文字、岗位关键词覆盖、经历证据强度、表达清晰度和 ATS 友好度等提示。匹配度是辅助参考，不是“能不能拿到面试”的保证。你要重点检查三件事：事实有没有被改错、数字是否真实、表达是否符合自己的说话习惯。')
    add_heading(doc, '5.2 继续追问和下载', 2)
    add_para(doc, '如果结果不合适，不需要重新上传。继续在底部输入框说明“哪里不对、要保留什么、想变成什么样”，系统会基于当前上下文继续改。确认后可下载结果，或进入“优化记录”保存和对比。')

    add_heading(doc, '六、优化记录：找回每一个岗位版本', 1)
    add_screenshot(doc, 'image5.png', '图3  优化记录示例（原版截图）', '记录页用于按公司、岗位、行业、模块和日期找回历史版本；页面数据仅作演示。')
    add_para(doc, '进入“优化记录”后，左侧是记录列表，右侧是当前选中版本的预览。你可以先按行业或模块筛选，再点击一条记录查看详细内容。')
    add_bullet(doc, '想继续修改：返回“求职工作台”，以当前记录为基础继续调整。')
    add_bullet(doc, '想对比不同版本：分别打开对应公司/岗位的记录，不要覆盖原版本。')
    add_bullet(doc, '想下载：在右侧预览区使用下载入口，下载前检查文件名和岗位名称。')
    add_callout(doc, '命名建议', '岗位名称尽量写清楚，例如“字节跳动｜高级前端工程师”，比“我的简历1”更容易在几周后找回来。')

    add_heading(doc, '七、面试助手：从会写到会说', 1)
    add_screenshot(doc, 'image6.png', '图4  面试助手示例（原版截图）', '面试助手目前包含面试题库、模拟面试、面试技巧和面试复盘四个页签。')
    add_table(doc, ['页签', '怎么用'], [
        ('面试题库', '按行业、岗位和难度筛题，先看题目，再整理自己的回答。'),
        ('模拟面试', '按题目逐个作答，练习把经历讲完整，而不是只背参考答案。'),
        ('面试技巧', '查看自我介绍、反问问题、线上面试环境等准备清单。'),
        ('面试复盘', '面试结束后记录公司、岗位、问题和自己的回答，再看改进建议。'),
    ], [1800, 7560])
    add_callout(doc, '更有效的练法', '先用自己的话回答，再看参考内容。把“我做过什么、遇到什么难点、最后结果怎样”说清楚，比背一段漂亮但不像自己的答案更重要。')

    add_heading(doc, '八、职业形象：把对外资料整理好', 1)
    add_screenshot(doc, 'image7.png', '图5  职业形象示例（原版截图）', '职业形象模块支持证件照风格生成，也支持名片等对外展示内容；生成前请确认照片和使用场景符合平台要求。')
    add_para(doc, '进入“职业形象”后，可以选择人物性别、参考风格和使用场景，上传清晰正面照后生成职业形象图。常见风格包括海马体证件风、商务金融风、科技简洁风、自然亲和风和创意品牌风；场景包括面试递交、内推沟通、线下社交和校招投递。')
    add_step(doc, 1, '上传清晰照片', '正脸、光线均匀、无遮挡的照片更容易生成自然结果。')
    add_step(doc, 2, '选择风格和场景', '按岗位和使用场景选择，不要为了“好看”选择与职业身份不匹配的风格。')
    add_step(doc, 3, '生成并检查', '检查脸部、服装、背景和文字是否自然；不满意时更换照片或风格重新生成。')

    add_heading(doc, '九、积分账单、会员额度和个人中心', 1)
    add_heading(doc, '9.1 积分账单', 2)
    add_para(doc, '“积分账单”用于回答两个问题：我现在还有多少积分？积分为什么增加或减少？进入后可查看余额、获取来源、消费用途和时间。遇到额度显示不一致时，先刷新页面并确认当前登录账号，再联系管理员核对。')
    add_heading(doc, '9.2 会员额度', 2)
    add_para(doc, '“会员额度”展示套餐、功能权益和开通申请。购买或申请前，先看清楚优化次数、分模块优化、高级模型、面试功能和历史记录保存等权益。当前版本支持提交开通申请，具体生效时间以管理员确认和页面状态为准。')
    add_heading(doc, '9.3 个人中心', 2)
    add_para(doc, '点击左侧头像或“个人中心”，可以查看账号资料、登录账号和基础设置。退出登录后，云端历史仍保存在账号中；未登录状态下可以体验部分页面，但不建议把重要简历只放在临时会话里。')

    add_heading(doc, '十、出错时怎么办', 1)
    add_table(doc, ['遇到的情况', '建议处理'], [
        ('登录失败', '检查账号是否输错、密码是否至少 6 位；确认是否切换到了正确的手机号/邮箱方式。'),
        ('短信验证码收不到', '确认手机号有效、网络正常，并等待倒计时结束后再重发；不要连续点击。'),
        ('PDF 解析为空', '优先换 DOCX/TXT，或直接粘贴简历正文；扫描图片 PDF 通常没有可读文本层。'),
        ('AI 结果太夸张', '明确要求“只使用真实经历，不新增事实和数字”，并逐段核对。'),
        ('结果不贴岗位', '重新检查 JD 是否完整，尤其是职责、技术栈和任职要求。'),
        ('页面没有更新', '刷新页面并确认登录账号；如果仍然异常，记录页面名称、操作步骤和报错文字。'),
    ], [2350, 7010])
    add_callout(doc, '重要原则', 'AI 生成内容一定要由本人最终确认。尤其是工作时间、项目数据、技术栈、证书和业绩数字，不能因为“看起来更专业”就直接使用。')

    add_heading(doc, '十一、完整使用示例：准备一份前端岗位简历', 1)
    add_step(doc, 1, '登录账号', '进入系统后使用手机号或邮箱登录；没有账号就先注册。')
    add_step(doc, 2, '新建岗位任务', '点击“求职工作台 → 新对话”，把目标岗位命名为“前端开发工程师”。')
    add_step(doc, 3, '粘贴 JD', '复制招聘页面的岗位描述，确认系统提取出了 React、TypeScript、权限、数据看板、性能优化等重点。')
    add_step(doc, 4, '上传简历', '上传 DOCX 或可复制文本的 PDF，检查工作经历和项目数据是否完整。')
    add_step(doc, 5, '先做整体诊断', '让系统先指出缺少的关键词和证据，再决定优先修改项目经历还是开场定位。')
    add_step(doc, 6, '分模块打磨', '输入“只优化项目经历，突出动作、难点和结果”，生成后逐条核对真实性。')
    add_step(doc, 7, '保存版本', '下载当前结果，并在“优化记录”里确认公司、岗位和日期，方便之后找回。')
    add_step(doc, 8, '准备面试', '到“面试助手”先做题库，再做模拟面试；面试后用“面试复盘”记录真实问题。')
    add_step(doc, 9, '补齐职业形象', '到“职业形象”整理证件照、名片和作品链接，统一对外展示风格。')

    add_heading(doc, '十二、功能入口速查', 1)
    add_table(doc, ['你现在想做的事', '入口', '下一步'], [
        ('改目标岗位或 JD', '求职工作台', '更新 JD / 粘贴 JD / 读取岗位链接'),
        ('上传或粘贴简历', '求职工作台', '上传文件或直接输入文本'),
        ('优化整份简历', '求职工作台', '选择整体诊断并发送要求'),
        ('只改项目/技能/自我介绍', '求职工作台', '选择对应模块再发送指令'),
        ('找回旧版本', '优化记录', '筛选并点击历史记录'),
        ('练习面试', '面试助手', '面试题库 / 模拟面试'),
        ('记录面试问题', '面试助手', '面试复盘'),
        ('生成证件照或整理名片', '职业形象', '选择风格、场景并生成'),
        ('看积分明细', '积分账单', '查看余额与流水'),
        ('看套餐或提交申请', '会员额度', '查看权益并提交开通申请'),
        ('改账号资料或退出', '个人中心', '查看资料、设置或退出登录'),
    ], [2700, 2200, 4460])
    add_para(doc, '本说明书以当前产品页面和代码中的功能入口为准。产品持续迭代时，建议每次发布新功能后同步更新版本号、截图和本手册中的入口名称。', 9.5, MUTED, italic=True, after=8)

    doc.core_properties.title = '职引简历AI求职工作台系统操作说明书'
    doc.core_properties.subject = '产品操作说明与使用指南'
    doc.core_properties.author = '职引简历'
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    build()
