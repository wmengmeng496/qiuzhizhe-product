from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = Path('/private/tmp/简历优化后端代码_完整摘录.docx')

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def add_code(doc, text):
    for line in text.splitlines() or ['']:
        p = doc.add_paragraph(style='Code')
        p.add_run(line if line else ' ')

def add_file(doc, rel, title=None, start=None, end=None):
    path = ROOT / rel
    content = path.read_text(encoding='utf-8')
    lines = content.splitlines()
    if start is not None:
        content = '\n'.join(lines[start-1:end])
        label = f'{rel}（第 {start}-{end or len(lines)} 行）'
    else:
        label = rel
    doc.add_heading(title or label, level=2)
    doc.add_paragraph(f'源码路径：{label}', style='Caption')
    add_code(doc, content)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
styles['Normal'].font.size = Pt(10)
for name, size, color in [('Title', 24, '17365D'), ('Heading 1', 16, '17365D'), ('Heading 2', 12, '2F5597')]:
    styles[name].font.name = 'Microsoft YaHei'
    styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)
if 'Code' not in styles:
    code = styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
else:
    code = styles['Code']
code.font.name = 'Menlo'
code._element.rPr.rFonts.set(qn('w:eastAsia'), '等线')
code.font.size = Pt(8)
code.font.color.rgb = RGBColor(40, 40, 40)
code.paragraph_format.space_after = Pt(0)
code.paragraph_format.line_spacing = 1.0

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('简历优化后端代码完整摘录')
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = RGBColor(23, 54, 93)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('项目：职引简历｜生成日期：2026-08-19').italic = True

doc.add_heading('文档说明', level=1)
doc.add_paragraph('本文件从当前仓库中复制与“简历优化”后端直接相关的代码，按功能模块组织。代码内容保持原样；每个代码块前标注仓库相对路径。前端页面代码未纳入本摘录。')
doc.add_paragraph('核心调用链：前端 POST /api/resume-optimizations → Express 路由 → JWT 与 Zod 校验 → 额度预扣 → AI 网关 → 相似度检查/重写 → PostgreSQL 保存 → 事件记录 → 返回结果。', style='Intense Quote')

doc.add_heading('一、AI 配置、模型与提示词', level=1)
add_file(doc, 'server/config/ai.config.js')

doc.add_heading('二、AI 调用与简历优化服务', level=1)
add_file(doc, 'server/src/ai.js', start=1, end=30)
add_file(doc, 'server/src/ai.js', title='AI JSON 调用、模型降级与解析', start=158, end=245)
add_file(doc, 'server/src/ai.js', title='简历优化函数与低差异检测', start=309, end=382)

doc.add_heading('三、简历优化 API 路由', level=1)
add_file(doc, 'server/src/routes.js', start=1, end=70)
add_file(doc, 'server/src/routes.js', title='POST /resume-optimizations', start=329, end=407)
add_file(doc, 'server/src/routes.js', title='GET /resume-optimizations', start=440, end=452)

doc.add_heading('四、额度、事务与退款', level=1)
add_file(doc, 'server/src/usage.js')
add_file(doc, 'server/src/db.js')

doc.add_heading('五、认证与当前用户', level=1)
add_file(doc, 'server/src/auth.js', title='requireAuth JWT 鉴权', start=168, end=207)
add_file(doc, 'server/src/config.js', title='套餐额度与模型权限', start=1, end=70)

doc.add_heading('六、数据库结构', level=1)
add_file(doc, 'database/schema.sql')

doc.add_heading('七、服务启动与自动建表', level=1)
add_file(doc, 'server/src/index.js')
add_file(doc, 'server/src/schema.js')

doc.add_heading('八、代码关系摘要', level=1)
summary = [
    ('入口', 'server/src/index.js', '挂载 /api 路由、解析 JSON、启动时执行数据库 schema。'),
    ('鉴权', 'server/src/auth.js', '读取 Bearer JWT，查询用户并注入 req.user。'),
    ('接口', 'server/src/routes.js', '校验请求、预扣额度、调用 AI、保存结果、返回 JSON。'),
    ('AI', 'server/src/ai.js', '兼容 Chat Completions，支持模型降级、JSON 解析和低差异重写。'),
    ('额度', 'server/src/usage.js', '按用户套餐和月份原子增加 resume_optimize_used，失败退款。'),
    ('持久化', 'database/schema.sql', '保存优化记录、额度和事件。'),
]
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
for cell, text in zip(table.rows[0].cells, ['模块', '文件', '职责']):
    cell.text = text
    set_cell_shading(cell, 'D9EAF7')
for row in summary:
    cells = table.add_row().cells
    for cell, text in zip(cells, row):
        cell.text = text

doc.add_paragraph('说明：本摘录只反映生成文档时工作区中的代码状态。部署时仍需配置 DATABASE_URL、JWT_SECRET、TOAPIS_API_KEY/CUSTOM_PROXY_KEY 等环境变量。', style='Intense Quote')
doc.save(OUT)
print(OUT)
