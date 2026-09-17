from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_BREAK

root = Path('/Users/wuqiong/Desktop/文件夹们/我自己的/求职者产品')
out = root/'求职者产品_代码整理.docx'
files = [root/'src/lib/pdf-reader.ts', root/'src/lib/docx-reader.ts', root/'src/pages/home/page.tsx']

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.7); sec.right_margin = Inches(.7)
normal = doc.styles['Normal']; normal.font.name = '等线'; normal.font.size = Pt(10)
title = doc.add_paragraph(); title.style = doc.styles['Title']
title.add_run('求职者产品｜业务代码整理')
doc.add_paragraph('包含 PDF 解析、DOCX 解析及首页业务页面源码。')

for idx, path in enumerate(files):
    h = doc.add_heading(path.name, level=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    for line in path.read_text(encoding='utf-8').splitlines():
        r = p.add_run(line + '\n')
        r.font.name = 'Menlo'; r.font.size = Pt(8)
    if idx < len(files)-1:
        doc.add_page_break()

doc.save(out)
print(out)
